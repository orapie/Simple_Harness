from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable


SUPPORTED_EXTENSIONS = {".txt", ".md", ".json", ".jsonl", ".csv", ".pdf", ".docx"}


@dataclass(frozen=True)
class Document:
    path: str
    text: str
    metadata: dict[str, str]


def load_documents(root: Path | str) -> list[Document]:
    data_root = Path(root).resolve()
    if not data_root.exists():
        raise FileNotFoundError(f"data root does not exist: {data_root}")
    if data_root.is_file():
        paths: Iterable[Path] = [data_root]
    else:
        paths = sorted(path for path in data_root.rglob("*") if path.is_file())

    documents: list[Document] = []
    for path in paths:
        if _should_skip(path):
            continue
        text = load_document_text(path)
        if text.strip():
            documents.append(
                Document(
                    path=str(path),
                    text=text,
                    metadata={"extension": path.suffix.lower(), "name": path.name},
                )
            )
    return documents


def load_document_text(path: Path | str) -> str:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return ""
    if suffix in {".txt", ".md"}:
        return file_path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".json":
        data = json.loads(file_path.read_text(encoding="utf-8"))
        return json.dumps(data, ensure_ascii=False, indent=2)
    if suffix == ".jsonl":
        return "\n".join(_format_jsonl(file_path))
    if suffix == ".csv":
        return "\n".join(_format_csv(file_path))
    if suffix == ".pdf":
        return _load_pdf(file_path)
    if suffix == ".docx":
        return _load_docx(file_path)
    return ""


def _format_jsonl(path: Path) -> list[str]:
    lines: list[str] = []
    for raw_line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        if not raw_line.strip():
            continue
        try:
            lines.append(json.dumps(json.loads(raw_line), ensure_ascii=False))
        except json.JSONDecodeError:
            lines.append(raw_line)
    return lines


def _format_csv(path: Path) -> list[str]:
    rows: list[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames:
            for row in reader:
                rows.append("；".join(f"{key}: {value}" for key, value in row.items()))
        else:
            handle.seek(0)
            for row in csv.reader(handle):
                rows.append("；".join(row))
    return rows


def _load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("reading PDF files requires: python -m pip install pypdf") from exc
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _load_docx(path: Path) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("reading DOCX files requires: python -m pip install python-docx") from exc
    document = DocxDocument(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def _should_skip(path: Path) -> bool:
    if any(part.startswith(".") for part in path.parts):
        return True
    return path.suffix.lower() not in SUPPORTED_EXTENSIONS
